"""
Google Sheets and Excel Exporter Module
Exports analyzed FTSE 100 stocks and ETF benchmarks to a professionally styled Excel workbook,
CSV files for 1-click Google Sheets import, and direct live publishing via gspread.
"""

import os
import pandas as pd
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class GoogleSheetsExporter:
    """
    Handles file export to Excel (.xlsx), CSV (.csv), and direct live sync to Google Sheets (gspread).
    """

    def __init__(self, output_dir: str = "."):
        self.output_dir = output_dir

    def format_excel_workbook(self, excel_path: str, df_stocks: pd.DataFrame, df_etfs: pd.DataFrame):
        """
        Creates a beautifully formatted multi-tab Excel workbook.
        """
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # Remove default sheet

        # Color Palette
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")  # Dark Navy
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        
        green_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")  # Soft Green
        dark_green_font = Font(name="Calibri", size=10, bold=True, color="155724")
        
        yellow_fill = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")  # Soft Yellow
        dark_yellow_font = Font(name="Calibri", size=10, bold=True, color="856404")
        
        red_fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")  # Soft Red
        dark_red_font = Font(name="Calibri", size=10, bold=True, color="721C24")

        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        # -------------------------------------------------------------
        # Tab 1: Executive Summary
        # -------------------------------------------------------------
        ws_sum = wb.create_sheet(title="Executive Summary")
        ws_sum.views.sheetView[0].showGridLines = True

        ws_sum["A1"] = "FTSE 100 Warren Buffett Value Investing Dashboard"
        ws_sum["A1"].font = Font(name="Calibri", size=16, bold=True, color="1F4E78")
        ws_sum["A2"] = "Quantitative Screening, Intrinsic Value (DCF), & ETF Benchmarks"
        ws_sum["A2"].font = Font(name="Calibri", size=11, italic=True, color="595959")

        # KPI Summary Cards
        total_count = len(df_stocks)
        strong_buys = len(df_stocks[df_stocks['Signal'] == 'STRONG BUY'])
        buys = len(df_stocks[df_stocks['Signal'] == 'BUY'])
        holds = len(df_stocks[df_stocks['Signal'] == 'HOLD'])
        sells = len(df_stocks[df_stocks['Signal'] == 'SELL'])

        kpis = [
            ("Total FTSE 100 Stocks", total_count, "1F4E78"),
            ("Strong Buy Signals", strong_buys, "155724"),
            ("Buy Signals", buys, "28A745"),
            ("Hold Signals", holds, "856404"),
            ("Sell / Overvalued Signals", sells, "721C24"),
        ]

        row = 4
        ws_sum.cell(row=row, column=1, value="Key Performance Indicators").font = Font(size=12, bold=True, color="1F4E78")
        row += 1

        ws_sum.cell(row=row, column=1, value="Metric").font = header_font
        ws_sum.cell(row=row, column=1).fill = header_fill
        ws_sum.cell(row=row, column=2, value="Count / Value").font = header_font
        ws_sum.cell(row=row, column=2).fill = header_fill

        for label, val, hex_col in kpis:
            row += 1
            c1 = ws_sum.cell(row=row, column=1, value=label)
            c2 = ws_sum.cell(row=row, column=2, value=val)
            c1.border = thin_border
            c2.border = thin_border
            c2.font = Font(bold=True, color=hex_col)
            c2.alignment = Alignment(horizontal="center")

        # Top 5 Value Stocks Table
        row += 3
        ws_sum.cell(row=row, column=1, value="Top 5 Warren Buffett Value Opportunities").font = Font(size=12, bold=True, color="1F4E78")
        row += 1

        top5 = df_stocks.sort_values(by='Buffett Score', ascending=False).head(5)
        top_cols = ['Ticker', 'Name', 'Sector', 'Buffett Score', 'Signal', 'Margin of Safety (%)', 'Price (£)', 'Intrinsic Value (£)']
        
        for c_idx, col_name in enumerate(top_cols, start=1):
            cell = ws_sum.cell(row=row, column=c_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for _, r_data in top5.iterrows():
            row += 1
            for c_idx, col_name in enumerate(top_cols, start=1):
                val = r_data[col_name]
                cell = ws_sum.cell(row=row, column=c_idx, value=val)
                cell.border = thin_border
                if col_name == 'Signal':
                    cell.alignment = Alignment(horizontal="center")
                    cell.font = dark_green_font
                    cell.fill = green_fill
                elif col_name in ['Buffett Score', 'Margin of Safety (%)']:
                    cell.alignment = Alignment(horizontal="right")
                    cell.font = Font(bold=True)

        # -------------------------------------------------------------
        # Tab 2: FTSE 100 Value Screener
        # -------------------------------------------------------------
        ws_screen = wb.create_sheet(title="FTSE 100 Screener")
        ws_screen.views.sheetView[0].showGridLines = True

        screener_cols = [
            'Ticker', 'Name', 'Sector', 'Signal', 'Buffett Score',
            'Price (£)', 'Intrinsic Value (£)', 'Margin of Safety (%)',
            'P/E Ratio', 'P/B Ratio', 'PEG Ratio', 'ROE (%)', 'Net Margin (%)',
            'Debt/Equity', 'Current Ratio', 'FCF Yield (%)', 'Dividend Yield (%)', 'Payout Ratio (%)'
        ]

        # Filter cols present in df_stocks
        display_cols = [c for c in screener_cols if c in df_stocks.columns]

        for c_idx, col_name in enumerate(display_cols, start=1):
            cell = ws_screen.cell(row=1, column=c_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for r_idx, r_data in df_stocks[display_cols].iterrows():
            row_num = r_idx + 2
            for c_idx, col_name in enumerate(display_cols, start=1):
                val = r_data[col_name]
                cell = ws_screen.cell(row=row_num, column=c_idx, value=val)
                cell.border = thin_border

                if col_name == 'Signal':
                    cell.alignment = Alignment(horizontal="center")
                    if val == 'STRONG BUY' or val == 'BUY':
                        cell.fill = green_fill
                        cell.font = dark_green_font
                    elif val == 'HOLD':
                        cell.fill = yellow_fill
                        cell.font = dark_yellow_font
                    else:
                        cell.fill = red_fill
                        cell.font = dark_red_font
                elif col_name == 'Buffett Score':
                    cell.alignment = Alignment(horizontal="right")
                    cell.font = Font(bold=True)

        # Auto-adjust column widths
        for ws in [ws_sum, ws_screen]:
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        # -------------------------------------------------------------
        # Tab 3: ETFs & Investment Vehicles
        # -------------------------------------------------------------
        ws_etf = wb.create_sheet(title="ETFs & Investment Vehicles")
        ws_etf.views.sheetView[0].showGridLines = True

        etf_cols = list(df_etfs.columns)
        for c_idx, col_name in enumerate(etf_cols, start=1):
            cell = ws_etf.cell(row=1, column=c_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for r_idx, r_data in df_etfs.iterrows():
            row_num = r_idx + 2
            for c_idx, col_name in enumerate(etf_cols, start=1):
                val = r_data[col_name]
                cell = ws_etf.cell(row=row_num, column=c_idx, value=val)
                cell.border = thin_border

        for col in ws_etf.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_etf.column_dimensions[col_letter].width = max(max_len + 3, 14)

        # Save workbook
        wb.save(excel_path)
        print(f"Successfully generated formatted Excel report: {excel_path}")

    def export_all(self, df_stocks: pd.DataFrame, df_etfs: pd.DataFrame) -> Dict[str, str]:
        """
        Exports data to Excel and CSV files.
        Returns paths to created files.
        """
        os.makedirs(self.output_dir, exist_ok=True)

        excel_path = os.path.join(self.output_dir, "ftse100_buffett_value_analysis.xlsx")
        csv_stocks_path = os.path.join(self.output_dir, "ftse100_screener.csv")
        csv_etf_path = os.path.join(self.output_dir, "etf_benchmark.csv")

        # Save CSV files for 1-click Google Sheets import
        df_stocks.to_csv(csv_stocks_path, index=False)
        df_etfs.to_csv(csv_etf_path, index=False)
        print(f"Exported CSV files:\n - {csv_stocks_path}\n - {csv_etf_path}")

        # Formatted Excel Export
        self.format_excel_workbook(excel_path, df_stocks, df_etfs)

        return {
            'excel': excel_path,
            'csv_stocks': csv_stocks_path,
            'csv_etfs': csv_etf_path
        }

    def publish_to_google_sheets(self, df_stocks: pd.DataFrame, df_etfs: pd.DataFrame, creds_path: Optional[str] = None) -> Optional[str]:
        """
        Optional live publishing to Google Sheets via gspread if service account credentials are provided.
        """
        try:
            import gspread
            from google.oauth2.service_account import Credentials

            if not creds_path or not os.path.exists(creds_path):
                print("No Google Service Account credentials provided/found. Skipping live Google Sheets publishing.")
                return None

            scopes = ['https://www.googleapis.com/auth/spreadsheets', 'https://www.googleapis.com/auth/drive']
            credentials = Credentials.from_service_account_file(creds_path, scopes=scopes)
            client = gspread.authorize(credentials)

            sheet_title = "FTSE 100 Warren Buffett Value Investing Screener"
            try:
                spreadsheet = client.open(sheet_title)
            except gspread.SpreadsheetNotFound:
                spreadsheet = client.create(sheet_title)

            # Update Stocks worksheet
            try:
                ws_stocks = spreadsheet.worksheet("FTSE 100 Screener")
            except gspread.WorksheetNotFound:
                ws_stocks = spreadsheet.add_worksheet(title="FTSE 100 Screener", rows="150", cols="25")

            ws_stocks.clear()
            ws_stocks.update([df_stocks.columns.values.tolist()] + df_stocks.fillna('').values.tolist())

            # Update ETF worksheet
            try:
                ws_etf = spreadsheet.worksheet("ETFs & Benchmarks")
            except gspread.WorksheetNotFound:
                ws_etf = spreadsheet.add_worksheet(title="ETFs & Benchmarks", rows="50", cols="20")

            ws_etf.clear()
            ws_etf.update([df_etfs.columns.values.tolist()] + df_etfs.fillna('').values.tolist())

            print(f"Successfully published live Google Sheet!\nURL: {spreadsheet.url}")
            return spreadsheet.url

        except Exception as e:
            print(f"Google Sheets publishing notice: {e}")
            return None

